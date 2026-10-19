"""Modern Office extractors (XLSX/XLSM, PPTX) and legacy XLS via xlrd (DOCX: word.py).

All open files read-only; macros are never executed (openpyxl and xlrd do not
run VBA; python-docx/pptx only parse XML).
"""

from __future__ import annotations

from pathlib import Path

from .base import Chunk, Extractor, ExtractResult, clean, col_letter, rows_to_markdown

ROW_BLOCK = 200


# ------------------------------------------------------------- sheet helper
def _sheet_chunks(name: str, rows: list[list], res: ExtractResult) -> None:
    """Split a sheet's rows into <=200-row markdown blocks with A1 ranges."""
    if not rows:
        return
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        return
    header = None
    for start in range(0, len(rows), ROW_BLOCK):
        block = rows[start:start + ROW_BLOCK]
        if not any(any(v not in (None, "") for v in r) for r in block):
            continue
        if header is None:
            header = block[0]
            body = block
        else:
            body = [header] + block   # repeat header so each chunk is self-describing
        md = rows_to_markdown(body)
        if not md:
            continue
        rng = f"{name}!A{start + 1}:{col_letter(width)}{start + len(block)}"
        res.chunks.append(Chunk(rng, md, has_table=True))


# --------------------------------------------------------------------- XLSX
class XlsxExtractor(Extractor):
    exts = ("xlsx", "xlsm")

    def extract(self, path: Path) -> ExtractResult:
        import openpyxl

        try:
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_vba=False)
        except Exception as exc:
            return ExtractResult(status="failed", error=f"open: {exc}")
        res = ExtractResult()
        try:
            p = wb.properties
            res.meta = {"title": (p.title or "").strip() or None,
                        "authors": (p.creator or "").strip() or None,
                        "year": p.created.year if p.created else None,
                        "sheets": wb.sheetnames}
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                _sheet_chunks(ws.title, rows, res)
        except Exception as exc:
            res.status, res.error = "failed", f"extract: {exc}"
        finally:
            wb.close()
        if res.status == "ok" and not res.chunks:
            res.status = "empty"
        return res


# ---------------------------------------------------------------- XLS (BIFF)
class XlsExtractor(Extractor):
    exts = ("xls",)

    def extract(self, path: Path) -> ExtractResult:
        import xlrd

        try:
            book = xlrd.open_workbook(str(path), on_demand=True)
        except Exception as exc:
            return ExtractResult(status="failed", error=f"open: {exc}")
        res = ExtractResult(meta={"title": None, "authors": getattr(book, "user_name", None) or None,
                                  "sheets": book.sheet_names()})
        try:
            for name in book.sheet_names():
                sh = book.sheet_by_name(name)
                rows = [sh.row_values(i) for i in range(sh.nrows)]
                _sheet_chunks(name, rows, res)
                book.unload_sheet(name)
        except Exception as exc:
            res.status, res.error = "failed", f"extract: {exc}"
        finally:
            book.release_resources()
        if res.status == "ok" and not res.chunks:
            res.status = "empty"
        return res


# --------------------------------------------------------------------- PPTX
class PptxExtractor(Extractor):
    exts = ("pptx",)

    def extract(self, path: Path) -> ExtractResult:
        from pptx import Presentation

        try:
            prs = Presentation(str(path))
        except Exception as exc:
            return ExtractResult(status="failed", error=f"open: {exc}")
        res = ExtractResult()
        cp = prs.core_properties
        res.meta = {"title": (cp.title or "").strip() or None,
                    "authors": (cp.author or "").strip() or None,
                    "year": cp.created.year if cp.created else None,
                    "slide_count": len(prs.slides)}
        for i, slide in enumerate(prs.slides, 1):
            parts, has_table = [], False
            for shp in slide.shapes:
                if getattr(shp, "has_text_frame", False) and shp.text_frame.text.strip():
                    parts.append(shp.text_frame.text)
                if getattr(shp, "has_table", False):
                    rows = [[c.text for c in r.cells] for r in shp.table.rows]
                    md = rows_to_markdown(rows)
                    if md:
                        parts.append(md)
                        has_table = True
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append("NOTES: " + notes)
            text = clean("\n".join(parts))
            if text:
                res.chunks.append(Chunk(f"slide {i}", text, has_table=has_table))
        if not res.meta["title"] and res.chunks:
            res.meta["title_guess"] = res.chunks[0].text.split("\n", 1)[0][:200]
        res.status = "ok" if res.chunks else "empty"
        return res
